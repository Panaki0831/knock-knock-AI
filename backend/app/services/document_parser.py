"""Document parser -- extracts text content from pptx, pdf, and docx files."""

from __future__ import annotations

import io
from dataclasses import dataclass

import structlog

logger = structlog.get_logger(__name__)


@dataclass
class ParseResult:
    """Result of parsing a document."""

    text: str
    page_count: int


def parse_pdf(file_bytes: bytes) -> ParseResult:
    """Extract text from a PDF file."""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(file_bytes))
    pages: list[str] = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            pages.append(text.strip())
    return ParseResult(text="\n\n".join(pages), page_count=len(reader.pages))


def parse_pptx(file_bytes: bytes) -> ParseResult:
    """Extract text from a PowerPoint file."""
    from pptx import Presentation

    prs = Presentation(io.BytesIO(file_bytes))
    slides_text: list[str] = []
    for slide_num, slide in enumerate(prs.slides, 1):
        parts: list[str] = [f"--- Slide {slide_num} ---"]
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                parts.append(shape.text.strip())
        slides_text.append("\n".join(parts))
    return ParseResult(text="\n\n".join(slides_text), page_count=len(prs.slides))


def parse_docx(file_bytes: bytes) -> ParseResult:
    """Extract text from a Word document."""
    from docx import Document

    doc = Document(io.BytesIO(file_bytes))
    paragraphs: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            paragraphs.append(para.text.strip())

    # Count approximate "pages" by paragraph blocks
    page_count = max(1, len(paragraphs) // 30)
    return ParseResult(text="\n\n".join(paragraphs), page_count=page_count)


def parse_document(file_bytes: bytes, filename: str) -> ParseResult:
    """Parse a document based on its file extension."""
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return parse_pdf(file_bytes)
    elif lower.endswith(".pptx"):
        return parse_pptx(file_bytes)
    elif lower.endswith(".docx"):
        return parse_docx(file_bytes)
    else:
        raise ValueError(f"Unsupported file type: {filename}")
